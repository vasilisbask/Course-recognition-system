from urllib.parse import urlencode
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.db.models import Case, IntegerField, Q, When
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.html import escape
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils import timezone
from django.utils.translation import get_language, gettext_lazy as _
from django.views.decorators.http import require_POST

from dal import autocomplete
from curricula.models import Course
from mailer.gmail import notify
from myprofile.checks import is_course_recognition_admin, is_course_recognition_professor, is_course_recognition_secretariat, is_course_recognition_student
from myprofile.models import StaffMember, Student
from scopes.models import Secretariat

from .forms import AdminCourseRecognitionRequestForm, CourseInstructorAssignmentForm, CourseInstructorImportForm, RecommendationForm, SecretariatProtocolForm, StudentCourseRecognitionRequestForm, courses_for_student
from .models import CourseRecognitionRequest
from .resources import CourseInstructorAssignmentResource, course_assignment_exists_message


def _student_for(user):
    student = Student.objects.filter(user=user).order_by("id").first()
    if not student:
        raise Http404
    return student


def _staff_for(user):
    return get_object_or_404(StaffMember, user=user)


def _student_full_name(student):
    return " ".join(part for part in (student.given_name, student.surname) if part)


def _professor_queryset(user):
    professor = _staff_for(user)
    # Professors can see only requests for courses assigned to them.
    return CourseRecognitionRequest.objects.filter(course__assigned_to=professor).distinct()


def _secretariat_programs(user):
    return (
        Secretariat.objects
        .filter(user=user, programs__isnull=False)
        .values_list("programs", flat=True)
        .distinct()
    )


def _secretariat_queryset(user):
    # Secretariat users work only inside the study programs linked to their account.
    return CourseRecognitionRequest.objects.filter(course__program_id__in=_secretariat_programs(user))


def _filter_course_autocomplete(queryset, query):
    queryset = queryset.filter(active=True)
    if query:
        queryset = queryset.filter(
            Q(title_gr__icontains=query)
            | Q(title_en__icontains=query)
            | Q(code_gr__icontains=query)
            | Q(code_en__icontains=query)
        )
    return queryset.order_by("semester", "title_gr")[:20]


class StudentCourseAutocomplete(LoginRequiredMixin, UserPassesTestMixin, autocomplete.Select2QuerySetView):
    def test_func(self):
        return is_course_recognition_student(self.request.user)

    def get_queryset(self):
        student = Student.objects.filter(user=self.request.user).order_by("id").first()
        return _filter_course_autocomplete(courses_for_student(student), self.q)


class SecretariatCourseAutocomplete(LoginRequiredMixin, UserPassesTestMixin, autocomplete.Select2QuerySetView):
    def test_func(self):
        return is_course_recognition_secretariat(self.request.user)

    def get_queryset(self):
        queryset = Course.objects.filter(program_id__in=_secretariat_programs(self.request.user))
        return _filter_course_autocomplete(queryset, self.q)


class AdminCourseAutocomplete(LoginRequiredMixin, UserPassesTestMixin, autocomplete.Select2QuerySetView):
    def test_func(self):
        return is_course_recognition_admin(self.request.user)

    def get_queryset(self):
        request_obj = get_object_or_404(CourseRecognitionRequest.objects.select_related("student"), pk=self.kwargs["pk"])
        student = request_obj.student
        return _filter_course_autocomplete(courses_for_student(student), self.q)


class ProfessorAutocomplete(LoginRequiredMixin, UserPassesTestMixin, autocomplete.Select2QuerySetView):
    def test_func(self):
        return is_course_recognition_secretariat(self.request.user)

    def get_queryset(self):
        queryset = StaffMember.objects.filter(user__is_active=True)
        if self.q:
            queryset = queryset.filter(
                Q(display_name__icontains=self.q)
                | Q(display_name_en__icontains=self.q)
                | Q(given_name__icontains=self.q)
                | Q(surname__icontains=self.q)
                | Q(email__icontains=self.q)
            )
        return queryset.order_by("surname", "given_name")[:20]


def _normalize_request_status_filter(value):
    valid_filters = {"all", "under_review", "completed", "inactive"}
    return value if value in valid_filters else "under_review"


def _apply_request_status_filter(queryset, status_filter):
    if status_filter == "under_review":
        return queryset.filter(status=CourseRecognitionRequest.UNDER_REVIEW)
    if status_filter == "completed":
        return queryset.filter(
            status__in=[
                CourseRecognitionRequest.APPROVED,
                CourseRecognitionRequest.REJECTED,
            ]
        )
    if status_filter == "inactive":
        return queryset.filter(status=CourseRecognitionRequest.WITHDRAWN)
    return queryset


def _email_list(values):
    return ",".join(sorted({value for value in values if value}))


def _split_email_list(value):
    if not value:
        return []
    if isinstance(value, str):
        return [email.strip() for email in value.split(",") if email.strip()]
    return [email for email in value if email]


def _merge_email_lists(*values):
    # Normalize comma-separated values before removing duplicate recipients.
    emails = []
    for value in values:
        emails.extend(_split_email_list(value))
    return _email_list(emails)


def _admin_emails():
    User = get_user_model()
    return _email_list(
        User.objects.filter(is_superuser=True, is_active=True)
        .exclude(email="")
        .values_list("email", flat=True)
    )


def _professor_emails(obj):
    if not obj.course_id:
        return ""
    return _email_list(obj.course.assigned_to.values_list("email", flat=True))


def _recommendation_professor_email(recommendation):
    if not recommendation:
        return ""
    if recommendation.professor_id:
        return recommendation.professor.email or (
            recommendation.professor.user and recommendation.professor.user.email
        ) or ""
    return ""


def _secretariat_emails(obj):
    if not obj.course_id or not obj.course.program_id:
        return ""

    program = obj.course.program
    query = Q(programs=program)
    if program.department_id:
        query |= Q(departments=program.department)
    secretariats = Secretariat.objects.filter(query).distinct()
    return _email_list(secretariats.values_list("user__email", flat=True))


def _send_request_notification(request, obj):
    recipient_emails = _merge_email_lists(
        _professor_emails(obj),
        _secretariat_emails(obj),
        _admin_emails(),
    )

    if not recipient_emails:
        return

    url = request.build_absolute_uri(reverse("myprofile:index"))
    student = escape(str(obj.student_display_name))
    course = escape(str(obj.course))
    original_course = escape(obj.title_original_course)

    body = f"""
    <p>Υποβλήθηκε νέα αίτηση αναγνώρισης μαθήματος.</p>
    <p>
        <strong>Φοιτητής:</strong> {student}<br>
        <strong>Μάθημα προς αναγνώριση:</strong> {course}<br>
        <strong>Περασμένο μάθημα:</strong> {original_course}
    </p>
    <p><a href="{url}">Μετάβαση στην εφαρμογή</a></p>
    """

    notify.delay(recipient_emails, "Νέα αίτηση αναγνώρισης μαθήματος", body)


def _student_email(obj):
    if not obj.student:
        return ""
    return obj.student.email or (obj.student.user and obj.student.user.email) or ""


def _send_professor_recommendation_notification(request, obj, recommendation):
    recipient_emails = _merge_email_lists(_secretariat_emails(obj), _admin_emails())
    if not recipient_emails:
        return

    url = request.build_absolute_uri(reverse("myprofile:index"))
    professor = escape(str(recommendation.professor or request.user))
    course = escape(str(obj.course))
    student = escape(str(obj.student_display_name))

    body = f"""
    <p>Ο καθηγητής υπέβαλε εισήγηση για αίτηση αναγνώρισης μαθήματος.</p>
    <p>
        <strong>Καθηγητής:</strong> {professor}<br>
        <strong>Φοιτητής:</strong> {student}<br>
        <strong>Μάθημα:</strong> {course}
    </p>
    <p>Η εισήγηση χρειάζεται αποδοχή ή απόρριψη.</p>
    <p><a href="{url}">Μετάβαση στην εφαρμογή</a></p>
    """

    notify.delay(recipient_emails, "Νέα εισήγηση καθηγητή για αίτηση αναγνώρισης μαθήματος", body)


def _send_professor_decision_notification(request, obj, recommendation=None, accepted=True):
    recipient_emails = _merge_email_lists(
        _recommendation_professor_email(recommendation),
        "" if recommendation else _professor_emails(obj),
    )
    if not recipient_emails:
        return

    url = request.build_absolute_uri(reverse("myprofile:index"))
    course = escape(str(obj.course))
    student = escape(str(obj.student_display_name))
    if recommendation:
        intro = "Η εισήγηση για την αίτηση αναγνώρισης μαθήματος έγινε αποδεκτή." if accepted else "Η εισήγηση για την αίτηση αναγνώρισης μαθήματος απορρίφθηκε."
    else:
        status_label = "εγκρίθηκε" if obj.status == CourseRecognitionRequest.APPROVED else "απορρίφθηκε"
        intro = f"Υπάρχει τελική απόφαση για την αίτηση αναγνώρισης μαθήματος: {status_label}."

    body = f"""
    <p>{intro}</p>
    <p>
        <strong>Φοιτητής:</strong> {student}<br>
        <strong>Μάθημα:</strong> {course}
    </p>
    <p><a href="{url}">Μετάβαση στην εφαρμογή</a></p>
    """

    notify.delay(recipient_emails, "Ενημέρωση εισήγησης αίτησης αναγνώρισης μαθήματος", body)


def _send_student_final_decision_notification(request, obj):
    student_email = _student_email(obj)
    if not student_email:
        return

    url = request.build_absolute_uri(reverse("myprofile:index"))
    course = escape(str(obj.course))
    status_label = "Εγκρίθηκε" if obj.status == CourseRecognitionRequest.APPROVED else "Απορρίφθηκε"
    decision_text = "εγκρίθηκε" if obj.status == CourseRecognitionRequest.APPROVED else "απορρίφθηκε"
    recommendation = getattr(obj, "recommendation", None)
    comments = ""
    if recommendation and recommendation.comments:
        comments = f'<br><strong>Σχόλια:</strong> {escape(recommendation.comments)}'

    body = f"""
    <p>Η αίτηση αναγνώρισης μαθήματος {decision_text}.</p>
    <p>
        <strong>Μάθημα:</strong> {course}<br>
        <strong>Αποτέλεσμα:</strong> {status_label}{comments}
    </p>
    <p><a href="{url}">Μετάβαση στην εφαρμογή</a></p>
    """

    notify.delay(student_email, "Απόφαση για αίτηση αναγνώρισης μαθήματος", body)


def _send_final_recommendation_notifications(request, obj, recommendation=None, notify_student=True):
    _send_professor_decision_notification(request, obj, recommendation=recommendation, accepted=True)
    if notify_student:
        _send_student_final_decision_notification(request, obj)


def _send_withdrawal_notification(request, obj):
    student_email = _student_email(obj)
    if not student_email:
        return

    url = request.build_absolute_uri(reverse("myprofile:index"))
    course = escape(str(obj.course))
    reason = escape(obj.withdrawal_reason) if obj.withdrawal_reason else ""

    body = f"""
    <p>Η αίτηση αναγνώρισης μαθήματος σημειώθηκε ως μη ενεργή από τη γραμματεία.</p>
    <p>
        <strong>Μάθημα:</strong> {course}<br>
        <strong>Αιτιολόγηση:</strong> {reason}
    </p>
    <p><a href="{url}">Μετάβαση στην εφαρμογή</a></p>
    """

    notify.delay(student_email, "Η αίτηση αναγνώρισης μαθήματος σημειώθηκε ως μη ενεργή", body)


def _send_admin_update_notification(request, obj):
    student_email = _student_email(obj)
    if not student_email:
        return

    url = request.build_absolute_uri(reverse("myprofile:index"))
    course = escape(str(obj.course))

    body = f"""
    <p>Η αίτηση αναγνώρισης μαθήματος ενημερώθηκε από τον διαχειριστή.</p>
    <p>
        <strong>Μάθημα:</strong> {course}
    </p>
    <p><a href="{url}">Μετάβαση στην εφαρμογή</a></p>
    """

    notify.delay(student_email, "Ενημέρωση αίτησης αναγνώρισης μαθήματος", body)


def _form_template(request, form, title, back_url, request_obj=None):
    return render(
        request,
        "course_recognition/form.html",
        {
            "form": form,
            "title": title,
            "back_url": back_url,
            "next_url": back_url,
            "request_obj": request_obj,
        },
    )


def _final_recommendation(obj):
    if obj.status not in (CourseRecognitionRequest.APPROVED, CourseRecognitionRequest.REJECTED):
        return None
    return getattr(obj, "recommendation", None)


def _pending_professor_recommendation(obj):
    # Professor recommendations stay pending until a secretariat or admin reviews them.
    if obj.status != CourseRecognitionRequest.UNDER_REVIEW:
        return None
    recommendation = getattr(obj, "recommendation", None)
    if recommendation and not recommendation.submitted_by_secretariat:
        return recommendation
    return None


def _add_missing_protocol_error(request, form):
    message = _("Δεν μπορεί να γίνει εισήγηση πριν καταχωρηθεί αριθμός πρωτοκόλλου.")
    form.add_error(None, message)


def _request_detail_template(request, obj, role, back_url, pending_recommendation=None):
    pending_professor_recommendation = _pending_professor_recommendation(obj)
    return render(
        request,
        "course_recognition/request_detail.html",
        {
            "request_obj": obj,
            "role": role,
            "back_url": back_url,
            "can_edit": (
                role == "student"
                and obj.status == CourseRecognitionRequest.UNDER_REVIEW
                and not pending_professor_recommendation
            ),
            "student_edit_locked_by_pending_recommendation": (
                role == "student" and bool(pending_professor_recommendation)
            ),
            "recommendation": _final_recommendation(obj),
            "pending_recommendation": pending_recommendation,
        },
    )


def _safe_next_url(request, fallback):
    next_url = request.POST.get("next") or request.GET.get("next")
    # Only redirect back to local URLs to avoid open redirect issues.
    if next_url and url_has_allowed_host_and_scheme(
        next_url,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return next_url
    return fallback


def _with_next(url, next_url):
    return f"{url}?{urlencode({'next': next_url})}"


def _select_request_for_actions(queryset):
    return queryset.select_related(
        "student",
        "student__program",
        "student__user",
        "course",
        "withdrawn_by",
        "recommendation",
        "recommendation__professor",
        "recommendation__submitted_by_user",
    )


@login_required
@user_passes_test(is_course_recognition_student)
def student_request_create(request):
    student = _student_for(request.user)
    if request.method == "POST":
        form = StudentCourseRecognitionRequestForm(request.POST, request.FILES, student=student)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.student = student
            obj.full_name = _student_full_name(student)
            obj.status = CourseRecognitionRequest.UNDER_REVIEW
            obj.save()
            _send_request_notification(request, obj)
            messages.success(request, _("Η αίτηση δημιουργήθηκε και είναι υπό εξέταση."))
            return redirect("course_recognition:student_request_detail", pk=obj.pk)
    else:
        form = StudentCourseRecognitionRequestForm(student=student)

    return _form_template(
        request,
        form,
        _("Νέα αίτηση αναγνώρισης"),
        reverse("myprofile:student_dashboard"),
    )


@login_required
@user_passes_test(is_course_recognition_student)
def student_request_detail(request, pk):
    student = _student_for(request.user)
    obj = get_object_or_404(
        CourseRecognitionRequest.objects.select_related(
            "student",
            "student__program",
            "student__user",
            "course",
            "withdrawn_by",
            "recommendation",
            "recommendation__professor",
            "recommendation__submitted_by_user",
        ),
        pk=pk,
        student=student,
    )
    return _request_detail_template(request, obj, "student", reverse("myprofile:student_dashboard"))


@login_required
@user_passes_test(is_course_recognition_student)
def student_request_update(request, pk):
    student = _student_for(request.user)
    obj = get_object_or_404(CourseRecognitionRequest, pk=pk, student=student)
    if obj.status != CourseRecognitionRequest.UNDER_REVIEW:
        messages.warning(
            request,
            _("Η αίτηση δεν μπορεί να επεξεργαστεί επειδή έχει ολοκληρωθεί η εισήγηση."),
        )
        return redirect("course_recognition:student_request_detail", pk=obj.pk)
    if _pending_professor_recommendation(obj):
        messages.warning(
            request,
            _("Η αίτηση δεν μπορεί να επεξεργαστεί επειδή υπάρχει εκκρεμής εισήγηση καθηγητή."),
        )
        return redirect("course_recognition:student_request_detail", pk=obj.pk)

    if request.method == "POST":
        form = StudentCourseRecognitionRequestForm(
            request.POST,
            request.FILES,
            instance=obj,
            student=student,
        )
        if form.is_valid():
            obj = form.save(commit=False)
            obj.student = student
            obj.full_name = _student_full_name(student)
            obj.save()
            messages.success(request, _("Η αίτηση ενημερώθηκε."))
            return redirect("course_recognition:student_request_detail", pk=obj.pk)
    else:
        form = StudentCourseRecognitionRequestForm(instance=obj, student=student)

    return _form_template(
        request,
        form,
        _("Επεξεργασία αίτησης"),
        reverse("course_recognition:student_request_detail", args=(obj.pk,)),
    )


@login_required
@user_passes_test(is_course_recognition_professor)
def professor_recommendation(request, pk):
    obj = get_object_or_404(
        _professor_queryset(request.user).select_related(
            "student",
            "student__program",
            "student__user",
            "course",
            "withdrawn_by",
            "recommendation",
            "recommendation__professor",
            "recommendation__submitted_by_user",
        ),
        pk=pk,
    )
    professor = _staff_for(request.user)
    recommendation = getattr(obj, "recommendation", None)
    if recommendation or obj.status == CourseRecognitionRequest.WITHDRAWN:
        return _request_detail_template(
            request,
            obj,
            "professor",
            reverse("myprofile:professor_dashboard"),
            pending_recommendation=_pending_professor_recommendation(obj),
        )

    if request.method == "POST":
        form = RecommendationForm(request.POST, instance=recommendation)
        # A protocol number is required before the request can move to recommendation.
        if not (obj.protocol_number or "").strip():
            _add_missing_protocol_error(request, form)
            return _form_template(
                request,
                form,
                _("Εισήγηση καθηγητή"),
                reverse("myprofile:professor_dashboard"),
                request_obj=obj,
            )
        if form.is_valid():
            recommendation = form.save(commit=False)
            recommendation.request = obj
            recommendation.professor = professor
            recommendation.submitted_by_secretariat = False
            recommendation.submitted_by_user = request.user
            recommendation.save()
            _send_professor_recommendation_notification(request, obj, recommendation)
            messages.success(request, _("Η εισήγηση αποθηκεύτηκε."))
            return redirect("course_recognition:professor_recommendation", pk=obj.pk)
    else:
        form = RecommendationForm(instance=recommendation)

    return _form_template(
        request,
        form,
        _("Εισήγηση καθηγητή"),
        reverse("myprofile:professor_dashboard"),
        request_obj=obj,
    )


@login_required
@user_passes_test(is_course_recognition_secretariat)
def secretariat_requests(request):
    base_queryset = (
        _secretariat_queryset(request.user)
        .select_related("student", "student__program", "student__user", "course", "recommendation")
        .order_by("-created_at")
    )
    status_filter = _normalize_request_status_filter(request.GET.get("status", "under_review"))
    queryset = _apply_request_status_filter(base_queryset, status_filter)

    status_filters = [
        {
            "key": "under_review",
            "label": _("Υπό εξέταση"),
            "count": base_queryset.filter(status=CourseRecognitionRequest.UNDER_REVIEW).count(),
        },
        {
            "key": "completed",
            "label": _("Ολοκληρωμένες"),
            "count": base_queryset.filter(
                status__in=[
                    CourseRecognitionRequest.APPROVED,
                    CourseRecognitionRequest.REJECTED,
                ]
            ).count(),
        },
        {
            "key": "inactive",
            "label": _("Μη ενεργές"),
            "count": base_queryset.filter(status=CourseRecognitionRequest.WITHDRAWN).count(),
        },
        {
            "key": "all",
            "label": _("Όλες"),
            "count": base_queryset.count(),
        },
    ]

    return render(
        request,
        "course_recognition/request_list.html",
        {
            "title": _("Αιτήσεις αναγνώρισης μαθημάτων"),
            "requests": queryset,
            "role": "secretariat",
            "status_filter": status_filter,
            "status_filters": status_filters,
        },
    )


def _import_course_assignments(user, uploaded_file):
    try:
        dataset = CourseInstructorAssignmentResource.dataset_from_upload(uploaded_file)
    except ValueError as exc:
        return 0, 0, [str(exc)], []

    resource = CourseInstructorAssignmentResource(program_ids=_secretariat_programs(user))
    return resource.import_assignments(dataset)


@login_required
@user_passes_test(is_course_recognition_secretariat)
def secretariat_course_assignment(request):
    import_form = CourseInstructorImportForm()
    is_english = (get_language() or "").startswith("en")

    if request.method == "POST" and "import_assignments" in request.POST:
        import_form = CourseInstructorImportForm(request.POST, request.FILES)
        form = CourseInstructorAssignmentForm(initial={"course": request.GET.get("course")}, user=request.user)
        if import_form.is_valid():
            imported_count, unchanged_count, import_errors, unchanged_messages = _import_course_assignments(
                request.user,
                import_form.cleaned_data["import_file"],
            )
            if import_errors:
                import_failed_prefix = (
                    ("The import completed with errors. " if imported_count else "The import was not completed. ")
                    if is_english
                    else (_("Το import ολοκληρώθηκε με σφάλματα. ") if imported_count else _("Το import δεν ολοκληρώθηκε. "))
                )
                messages.error(request, str(import_failed_prefix) + " ".join(str(error) for error in import_errors))
            for unchanged_message in unchanged_messages:
                messages.warning(request, unchanged_message)
            if imported_count:
                if imported_count == 1:
                    import_success = "Added 1 assignment." if is_english else _("Προστέθηκε 1 ανάθεση.")
                else:
                    import_success = (
                        "Added %(imported)s assignments."
                        if is_english
                        else _("Προστέθηκαν %(imported)s αναθέσεις.")
                    ) % {"imported": imported_count}
                messages.success(request, import_success)
            return redirect("course_recognition:secretariat_course_assignment")
    elif request.method == "POST":
        form = CourseInstructorAssignmentForm(request.POST, user=request.user)
        if form.is_valid():
            course, instructor, created = form.save()
            if created:
                messages.success(request, _("Η ανάθεση καθηγητή στο μάθημα αποθηκεύτηκε."))
            else:
                messages.warning(
                    request,
                    course_assignment_exists_message(instructor, course),
                )
            return redirect("course_recognition:secretariat_course_assignment")
    else:
        form = CourseInstructorAssignmentForm(initial={"course": request.GET.get("course")}, user=request.user)

    return render(
        request,
        "course_recognition/course_assignment.html",
        {
            "title": _("Ανάθεση μαθημάτων σε καθηγητές"),
            "form": form,
            "import_form": import_form,
            "back_url": reverse("myprofile:secretariat_dashboard"),
        },
    )


@login_required
@user_passes_test(is_course_recognition_secretariat)
def secretariat_protocol(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:secretariat_requests"))
    obj = get_object_or_404(
        _secretariat_queryset(request.user).select_related(
            "student",
            "student__program",
            "student__user",
            "course",
            "withdrawn_by",
            "recommendation",
            "recommendation__professor",
            "recommendation__submitted_by_user",
        ),
        pk=pk,
    )
    if _final_recommendation(obj) or obj.status == CourseRecognitionRequest.WITHDRAWN or obj.protocol_number:
        messages.warning(
            request,
            _("Ο αριθμός πρωτοκόλλου δεν μπορεί να αλλάξει για αυτή την αίτηση."),
        )
        return redirect(_with_next(reverse("course_recognition:secretariat_recommendation", args=(obj.pk,)), back_url))

    if request.method == "POST":
        form = SecretariatProtocolForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            messages.success(request, _("Ο αριθμός πρωτοκόλλου αποθηκεύτηκε."))
            return redirect(back_url)
    else:
        form = SecretariatProtocolForm(instance=obj)

    return _form_template(
        request,
        form,
        _("Αριθμός πρωτοκόλλου"),
        back_url,
        request_obj=obj,
    )


@login_required
@user_passes_test(is_course_recognition_secretariat)
def secretariat_recommendation(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:secretariat_requests"))
    obj = get_object_or_404(
        _secretariat_queryset(request.user).select_related(
            "student",
            "student__program",
            "student__user",
            "course",
            "withdrawn_by",
            "recommendation",
            "recommendation__professor",
            "recommendation__submitted_by_user",
        ),
        pk=pk,
    )
    recommendation = getattr(obj, "recommendation", None)
    if _final_recommendation(obj) or obj.status == CourseRecognitionRequest.WITHDRAWN:
        return _request_detail_template(request, obj, "secretariat", back_url)
    pending_recommendation = _pending_professor_recommendation(obj)
    if pending_recommendation:
        # Pending professor recommendations must be accepted or rejected first.
        return _request_detail_template(
            request,
            obj,
            "secretariat",
            back_url,
            pending_recommendation=pending_recommendation,
        )

    if request.method == "POST":
        form = RecommendationForm(request.POST, instance=recommendation)
        if not (obj.protocol_number or "").strip():
            _add_missing_protocol_error(request, form)
            return _form_template(
                request,
                form,
                _("Εισήγηση γραμματείας"),
                back_url,
                request_obj=obj,
            )
        if form.is_valid():
            recommendation = form.save(commit=False)
            recommendation.request = obj
            recommendation.professor = None
            recommendation.submitted_by_secretariat = True
            recommendation.submitted_by_user = request.user
            recommendation.save()
            # Secretariat recommendations are final decisions for the request.
            obj.status = (
                CourseRecognitionRequest.APPROVED
                if recommendation.approval
                else CourseRecognitionRequest.REJECTED
            )
            obj.save(update_fields=["status", "updated_at"])
            _send_final_recommendation_notifications(request, obj, recommendation=recommendation)
            messages.success(request, _("Η εισήγηση αποθηκεύτηκε."))
            return redirect(back_url)
    else:
        form = RecommendationForm(instance=recommendation)

    return _form_template(
        request,
        form,
        _("Εισήγηση γραμματείας"),
        back_url,
        request_obj=obj,
    )


def _accept_professor_recommendation(request, obj, back_url):
    recommendation = _pending_professor_recommendation(obj)
    if not recommendation:
        messages.warning(request, _("Δεν υπάρχει εκκρεμής εισήγηση καθηγητή για αποδοχή."))
        return redirect(back_url)

    # Accepting the professor recommendation turns it into the final decision.
    obj.status = (
        CourseRecognitionRequest.APPROVED
        if recommendation.approval
        else CourseRecognitionRequest.REJECTED
    )
    obj.save(update_fields=["status", "updated_at"])
    _send_final_recommendation_notifications(request, obj, recommendation=recommendation)
    messages.success(request, _("Η εισήγηση του καθηγητή έγινε αποδεκτή."))
    return redirect(back_url)


def _reject_professor_recommendation(request, obj, back_url):
    recommendation = _pending_professor_recommendation(obj)
    if not recommendation:
        messages.warning(request, _("Δεν υπάρχει εκκρεμής εισήγηση καθηγητή για απόρριψη."))
        return redirect(back_url)

    _send_professor_decision_notification(request, obj, recommendation=recommendation, accepted=False)
    # Rejection removes the pending recommendation so the request can be handled again.
    recommendation.delete()
    obj.status = CourseRecognitionRequest.UNDER_REVIEW
    obj.save(update_fields=["status", "updated_at"])
    messages.success(request, _("Η εισήγηση του καθηγητή απορρίφθηκε και η αίτηση παραμένει υπό εξέταση."))
    return redirect(back_url)


@login_required
@user_passes_test(is_course_recognition_secretariat)
@require_POST
def secretariat_accept_recommendation(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:secretariat_requests"))
    obj = get_object_or_404(_select_request_for_actions(_secretariat_queryset(request.user)), pk=pk)
    return _accept_professor_recommendation(request, obj, back_url)


@login_required
@user_passes_test(is_course_recognition_secretariat)
@require_POST
def secretariat_reject_recommendation(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:secretariat_requests"))
    obj = get_object_or_404(_select_request_for_actions(_secretariat_queryset(request.user)), pk=pk)
    return _reject_professor_recommendation(request, obj, back_url)


@login_required
@user_passes_test(is_course_recognition_secretariat)
@require_POST
def secretariat_withdraw(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:secretariat_requests"))
    obj = get_object_or_404(_secretariat_queryset(request.user), pk=pk)
    if obj.status != CourseRecognitionRequest.UNDER_REVIEW:
        messages.warning(request, _("Μόνο αιτήσεις που είναι υπό εξέταση μπορούν να σημειωθούν ως μη ενεργές."))
        return redirect(back_url)

    withdrawal_reason = request.POST.get("withdrawal_reason", "").strip()
    if not withdrawal_reason:
        messages.error(request, _("Η αιτιολόγηση διαγραφής είναι υποχρεωτική."))
        return redirect(back_url)

    # Withdrawal is a soft delete; the request remains available for audit and restore.
    obj.status = CourseRecognitionRequest.WITHDRAWN
    obj.withdrawal_reason = withdrawal_reason
    obj.withdrawn_by = request.user
    obj.withdrawn_at = timezone.now()
    obj.save(update_fields=["status", "withdrawal_reason", "withdrawn_by", "withdrawn_at", "updated_at"])
    _send_withdrawal_notification(request, obj)
    messages.success(request, _("Η αίτηση σημειώθηκε ως μη ενεργή."))
    return redirect(back_url)


@login_required
@user_passes_test(is_course_recognition_secretariat)
@require_POST
def secretariat_restore(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:secretariat_requests"))
    obj = get_object_or_404(_secretariat_queryset(request.user), pk=pk)
    if obj.status != CourseRecognitionRequest.WITHDRAWN:
        messages.warning(request, _("Μόνο μη ενεργές αιτήσεις μπορούν να επανέλθουν."))
        return redirect(back_url)

    obj.status = CourseRecognitionRequest.UNDER_REVIEW
    obj.withdrawal_reason = ""
    obj.withdrawn_by = None
    obj.withdrawn_at = None
    obj.save(update_fields=["status", "withdrawal_reason", "withdrawn_by", "withdrawn_at", "updated_at"])
    messages.success(request, _("Η αίτηση επανήλθε σε κατάσταση υπό εξέταση."))
    return redirect(back_url)


def _apply_export_visible_ids(request, queryset):
    ids_param = request.GET.get("ids")
    if ids_param is None:
        return queryset.order_by("id")

    # Export only the rows currently visible in the DataTables page.
    ids = []
    for raw_id in ids_param.split(","):
        raw_id = raw_id.strip()
        if not raw_id:
            continue
        try:
            ids.append(int(raw_id))
        except ValueError:
            continue

    if not ids:
        return queryset.none()

    ordering = Case(
        *[When(pk=pk, then=position) for position, pk in enumerate(ids)],
        output_field=IntegerField(),
    )
    return queryset.filter(pk__in=ids).order_by(ordering)


@login_required
@user_passes_test(is_course_recognition_secretariat)
def secretariat_export_xlsx(request):
    status_filter = _normalize_request_status_filter(request.GET.get("status", "under_review"))
    queryset = (
        _apply_request_status_filter(_secretariat_queryset(request.user), status_filter)
        .select_related("student", "course")
    )
    return _requests_xlsx_response(_apply_export_visible_ids(request, queryset))


def _requests_xlsx_response(queryset):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    is_english = (get_language() or "").startswith("en")

    wb = Workbook()
    ws = wb.active
    ws.title = "Recognition Requests" if is_english else "Αιτήσεις Αναγνώρισης"

    # Columns
    if is_english:
        headers = [
            "ID",
            "Full name",
            "Registration number",
            "Course to recognize",
            "Previously completed course",
            "Institution",
            "School",
            "Department",
            "Instructors",
            "Academic year",
            "Academic semester",
            "Theory grade",
            "Lab grade",
            "Average grade",
            "Theory hours",
            "Lab hours",
            "Protocol number",
            "Status",
            "Recommendation",
            "Recognition grade",
            "Recommendation comments",
        ]
        status_labels = {
            CourseRecognitionRequest.UNDER_REVIEW: "Under review",
            CourseRecognitionRequest.APPROVED: "Approved",
            CourseRecognitionRequest.REJECTED: "Rejected",
            CourseRecognitionRequest.WITHDRAWN: "Inactive",
        }
        semester_labels = {
            CourseRecognitionRequest.WINTER: "Winter",
            CourseRecognitionRequest.SPRING: "Spring",
        }
        year_labels = {
            CourseRecognitionRequest.YEAR_1: "1st",
            CourseRecognitionRequest.YEAR_2: "2nd",
            CourseRecognitionRequest.YEAR_3: "3rd",
            CourseRecognitionRequest.YEAR_4: "4th",
            CourseRecognitionRequest.YEAR_5: "5th",
        }
        approved_label = "Approved"
        rejected_label = "Rejected"
    else:
        headers = [
            "ID",
            "Ονοματεπώνυμο",
            "Αριθμός μητρώου",
            "Μάθημα προς αναγνώριση",
            "Περασμένο μάθημα",
            "Ίδρυμα",
            "Σχολή",
            "Τμήμα",
            "Διδάσκοντες",
            "Διδακτικό έτος",
            "Διδακτικό εξάμηνο",
            "Βαθμός θεωρίας",
            "Βαθμός εργαστηρίου",
            "Μέσος όρος",
            "Ώρες θεωρίας",
            "Ώρες εργαστηρίου",
            "Αριθμός πρωτοκόλλου",
            "Κατάσταση",
            "Εισήγηση",
            "Βαθμός αναγνώρισης",
            "Σχόλια εισήγησης",
        ]
        status_labels = {}
        semester_labels = {}
        year_labels = {}
        approved_label = str(_("Εγκρίθηκε"))
        rejected_label = str(_("Απορρίφθηκε"))

    ws.append(headers)

    # Style Header Row
    header_font = Font(name="Calibri", size=11, bold=True)
    header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for obj in queryset:
        recommendation = getattr(obj, "recommendation", None)
        recommendation_text = ""
        recognition_grade = ""
        if recommendation:
            if recommendation.approval:
                recommendation_text = approved_label
                recognition_grade = recommendation.grade or ""
            else:
                recommendation_text = rejected_label

        course_title = ""
        if obj.course:
            course_title = obj.course.title_en if is_english and obj.course.title_en else str(obj.course)

        row = [
            obj.id,
            obj.student_display_name,
            obj.student.reg_num if obj.student else "",
            course_title,
            obj.title_original_course,
            obj.institution,
            obj.school,
            obj.department,
            obj.instructors,
            year_labels.get(obj.academic_year, str(obj.get_academic_year_display() or "")),
            semester_labels.get(obj.academic_semester, str(obj.get_academic_semester_display() or "")),
            obj.grade_theory if obj.grade_theory is not None else "",
            obj.grade_lab if obj.grade_lab is not None else "",
            obj.average_grade if obj.average_grade is not None else "",
            obj.hours_theory if obj.hours_theory is not None else "",
            obj.hours_lab if obj.hours_lab is not None else "",
            obj.protocol_number,
            status_labels.get(obj.status, str(obj.get_status_display() or "")),
            recommendation_text,
            recognition_grade,
            recommendation.comments if recommendation else "",
        ]
        ws.append(row)

    # Apply formatting: auto-adjust column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            lines = val_str.split("\n")
            for line in lines:
                if len(line) > max_len:
                    max_len = len(line)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

    # Set row height for headers
    ws.row_dimensions[1].height = 28

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="course-recognition-requests.xlsx"'
    wb.save(response)

    return response


@login_required
@user_passes_test(is_course_recognition_admin)
def admin_export_xlsx(request):
    status_filter = _normalize_request_status_filter(request.GET.get("status", "under_review"))
    queryset = (
        _apply_request_status_filter(CourseRecognitionRequest.objects.all(), status_filter)
        .select_related("student", "course")
    )
    return _requests_xlsx_response(_apply_export_visible_ids(request, queryset))


@login_required
@user_passes_test(is_course_recognition_admin)
def admin_requests(request):
    base_queryset = (
        CourseRecognitionRequest.objects
        .select_related("student", "student__program", "student__user", "course", "recommendation")
        .order_by("-created_at")
    )
    status_filter = _normalize_request_status_filter(request.GET.get("status", "under_review"))
    queryset = _apply_request_status_filter(base_queryset, status_filter)
    status_filters = [
        {
            "key": "under_review",
            "label": _("Υπό εξέταση"),
            "count": base_queryset.filter(status=CourseRecognitionRequest.UNDER_REVIEW).count(),
        },
        {
            "key": "completed",
            "label": _("Ολοκληρωμένες"),
            "count": base_queryset.filter(
                status__in=[
                    CourseRecognitionRequest.APPROVED,
                    CourseRecognitionRequest.REJECTED,
                ]
            ).count(),
        },
        {
            "key": "inactive",
            "label": _("Μη ενεργές"),
            "count": base_queryset.filter(status=CourseRecognitionRequest.WITHDRAWN).count(),
        },
        {
            "key": "all",
            "label": _("Όλες"),
            "count": base_queryset.count(),
        },
    ]
    return render(
        request,
        "course_recognition/admin_request_list.html",
        {
            "title": _("Προβολή Αιτήσεων"),
            "requests": queryset,
            "status_filter": status_filter,
            "status_filters": status_filters,
        },
    )


@login_required
@user_passes_test(is_course_recognition_admin)
def admin_request_detail(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:admin_requests"))
    obj = get_object_or_404(
        CourseRecognitionRequest.objects.select_related(
            "student",
            "student__program",
            "student__user",
            "course",
            "withdrawn_by",
            "recommendation",
            "recommendation__professor",
            "recommendation__submitted_by_user",
        ),
        pk=pk,
    )
    return _request_detail_template(
        request,
        obj,
        "admin",
        back_url,
        pending_recommendation=_pending_professor_recommendation(obj),
    )


@login_required
@user_passes_test(is_course_recognition_admin)
def admin_request_update(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:admin_requests"))
    obj = get_object_or_404(CourseRecognitionRequest, pk=pk)
    if _final_recommendation(obj) or obj.status == CourseRecognitionRequest.WITHDRAWN:
        messages.warning(
            request,
            _("Η αίτηση δεν μπορεί να επεξεργαστεί επειδή δεν είναι υπό εξέταση."),
        )
        return redirect(_with_next(reverse("course_recognition:admin_request_detail", args=(obj.pk,)), back_url))
    if _pending_professor_recommendation(obj):
        messages.warning(
            request,
            _("Η αίτηση δεν μπορεί να επεξεργαστεί επειδή υπάρχει εκκρεμής εισήγηση καθηγητή."),
        )
        return redirect(_with_next(reverse("course_recognition:admin_request_detail", args=(obj.pk,)), back_url))

    if request.method == "POST":
        form = AdminCourseRecognitionRequestForm(request.POST, request.FILES, instance=obj)
        if form.is_valid():
            form.save()
            _send_admin_update_notification(request, obj)
            messages.success(request, _("Η αίτηση αποθηκεύτηκε."))
            return redirect(back_url)
    else:
        form = AdminCourseRecognitionRequestForm(instance=obj)

    return _form_template(
        request,
        form,
        _("Επεξεργασία αίτησης"),
        back_url,
    )


@login_required
@user_passes_test(is_course_recognition_admin)
def admin_protocol(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:admin_requests"))
    obj = get_object_or_404(_select_request_for_actions(CourseRecognitionRequest.objects.all()), pk=pk)
    if _final_recommendation(obj) or obj.status == CourseRecognitionRequest.WITHDRAWN or obj.protocol_number:
        messages.warning(
            request,
            _("Ο αριθμός πρωτοκόλλου δεν μπορεί να αλλάξει για αυτή την αίτηση."),
        )
        return redirect(_with_next(reverse("course_recognition:admin_request_detail", args=(obj.pk,)), back_url))

    if request.method == "POST":
        form = SecretariatProtocolForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            messages.success(request, _("Ο αριθμός πρωτοκόλλου αποθηκεύτηκε."))
            return redirect(back_url)
    else:
        form = SecretariatProtocolForm(instance=obj)

    return _form_template(
        request,
        form,
        _("Αριθμός πρωτοκόλλου"),
        back_url,
        request_obj=obj,
    )


@login_required
@user_passes_test(is_course_recognition_admin)
def admin_recommendation(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:admin_requests"))
    obj = get_object_or_404(_select_request_for_actions(CourseRecognitionRequest.objects.all()), pk=pk)
    recommendation = getattr(obj, "recommendation", None)
    if _final_recommendation(obj) or obj.status == CourseRecognitionRequest.WITHDRAWN:
        return _request_detail_template(request, obj, "admin", back_url)
    pending_recommendation = _pending_professor_recommendation(obj)
    if pending_recommendation:
        return _request_detail_template(
            request,
            obj,
            "admin",
            back_url,
            pending_recommendation=pending_recommendation,
        )

    if request.method == "POST":
        form = RecommendationForm(request.POST, instance=recommendation)
        if not (obj.protocol_number or "").strip():
            _add_missing_protocol_error(request, form)
            return _form_template(
                request,
                form,
                _("Εισήγηση"),
                back_url,
                request_obj=obj,
            )
        if form.is_valid():
            recommendation = form.save(commit=False)
            recommendation.request = obj
            recommendation.professor = None
            recommendation.submitted_by_secretariat = True
            recommendation.submitted_by_user = request.user
            recommendation.save()
            obj.status = (
                CourseRecognitionRequest.APPROVED
                if recommendation.approval
                else CourseRecognitionRequest.REJECTED
            )
            obj.save(update_fields=["status", "updated_at"])
            _send_final_recommendation_notifications(request, obj, recommendation=recommendation)
            messages.success(request, _("Η εισήγηση αποθηκεύτηκε."))
            return redirect(back_url)
    else:
        form = RecommendationForm(instance=recommendation)

    return _form_template(
        request,
        form,
        _("Εισήγηση"),
        back_url,
        request_obj=obj,
    )


@login_required
@user_passes_test(is_course_recognition_admin)
@require_POST
def admin_accept_recommendation(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:admin_requests"))
    obj = get_object_or_404(_select_request_for_actions(CourseRecognitionRequest.objects.all()), pk=pk)
    return _accept_professor_recommendation(request, obj, back_url)


@login_required
@user_passes_test(is_course_recognition_admin)
@require_POST
def admin_reject_recommendation(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:admin_requests"))
    obj = get_object_or_404(_select_request_for_actions(CourseRecognitionRequest.objects.all()), pk=pk)
    return _reject_professor_recommendation(request, obj, back_url)


@login_required
@user_passes_test(is_course_recognition_admin)
@require_POST
def admin_withdraw(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:admin_requests"))
    obj = get_object_or_404(CourseRecognitionRequest.objects.all(), pk=pk)
    if obj.status != CourseRecognitionRequest.UNDER_REVIEW:
        messages.warning(request, _("Μόνο αιτήσεις που είναι υπό εξέταση μπορούν να σημειωθούν ως μη ενεργές."))
        return redirect(back_url)

    withdrawal_reason = request.POST.get("withdrawal_reason", "").strip()
    if not withdrawal_reason:
        messages.error(request, _("Η αιτιολόγηση διαγραφής είναι υποχρεωτική."))
        return redirect(back_url)

    obj.status = CourseRecognitionRequest.WITHDRAWN
    obj.withdrawal_reason = withdrawal_reason
    obj.withdrawn_by = request.user
    obj.withdrawn_at = timezone.now()
    obj.save(update_fields=["status", "withdrawal_reason", "withdrawn_by", "withdrawn_at", "updated_at"])
    _send_withdrawal_notification(request, obj)
    messages.success(request, _("Η αίτηση σημειώθηκε ως μη ενεργή."))
    return redirect(back_url)


@login_required
@user_passes_test(is_course_recognition_admin)
@require_POST
def admin_restore(request, pk):
    back_url = _safe_next_url(request, reverse("course_recognition:admin_requests"))
    obj = get_object_or_404(CourseRecognitionRequest.objects.all(), pk=pk)
    if obj.status != CourseRecognitionRequest.WITHDRAWN:
        messages.warning(request, _("Μόνο μη ενεργές αιτήσεις μπορούν να επανέλθουν."))
        return redirect(back_url)

    obj.status = CourseRecognitionRequest.UNDER_REVIEW
    obj.withdrawal_reason = ""
    obj.withdrawn_by = None
    obj.withdrawn_at = None
    obj.save(update_fields=["status", "withdrawal_reason", "withdrawn_by", "withdrawn_at", "updated_at"])
    messages.success(request, _("Η αίτηση επανήλθε σε κατάσταση υπό εξέταση."))
    return redirect(back_url)
